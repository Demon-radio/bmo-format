"""Oracle conformance: REAL libraries verify OUR work both ways.
- Our exports must parse in python-docx/openpyxl/python-pptx with equal content.
- Our imports must agree with the real libs on genuine Office files.
Run: python tests/converters/test_oracles.py
"""
import sys
from io import BytesIO
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent / "converters" / "python"))
from bmo_ooxml import docx, pptx, xlsx  # noqa: E402

passed = failed = 0

def check(name, fn):
    global passed, failed
    try:
        fn()
        print(f"  PASS {name}")
        passed += 1
    except Exception as e:
        print(f"  FAIL {name}: {type(e).__name__}: {e}")
        failed += 1

def t_export_docx_readable():
    import docx as real
    title, envs, _, _ = docx.import_docx(docx.sample_docx())
    approx = []
    out = docx.export_docx(envs, title, approx)
    assert approx == []
    doc = real.Document(BytesIO(out))
    texts = ["".join(r.text for r in p.runs) for p in doc.paragraphs]
    assert "Annual Report" in texts and "Hello BMO" in "".join(texts), texts
    assert any("First item" in t for t in texts)
    tbl = doc.tables[0]
    got = [[c.text for c in row.cells] for row in tbl.rows]
    assert got == [["R1C1", "R1C2"], ["R2C1", "R2C2"]], got

def t_export_xlsx_readable():
    import openpyxl
    title, envs, _, _ = xlsx.import_xlsx(xlsx.sample_xlsx())
    approx = []
    out = xlsx.export_xlsx(envs, title, approx)
    wb = openpyxl.load_workbook(BytesIO(out))
    assert wb.sheetnames == ["Summary", "Q1"], wb.sheetnames
    ws = wb["Summary"]
    assert ws["A1"].value == 42
    # B1 is covered by the A1:B1 merge: Excel shows A1's value there, so
    # openpyxl reports None (correct merged-cell semantics; B1's own stored
    # value round-trips through OUR pipeline — see test_converters xlsx 100%).
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:B1"]
    assert ws["D1"].value == "=A1*2"  # formula text preserved (openpyxl adds =)
    assert ws["C1"].number_format == "0.00%"
    assert ws["A2"].value is True and ws["B2"].value == "#DIV/0!"

def t_export_pptx_readable():
    from pptx import Presentation
    title, envs, _, _ = pptx.import_pptx(pptx.sample_pptx())
    approx = []
    out = pptx.export_pptx(envs, title, approx, {})
    prs = Presentation(BytesIO(out))
    assert len(prs.slides) == 2
    texts = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                texts.append(sh.text)
    assert any("Hello BMO" in t for t in texts), texts
    assert any("Second" in t for t in texts), texts

def t_differential_docx():
    import docx as real
    data = (HERE / "samples" / "real-images.docx").read_bytes()
    doc = real.Document(BytesIO(data))
    real_texts = sorted("".join(r.text for r in p.runs) for p in doc.paragraphs
                        if "".join(r.text for r in p.runs).strip())
    _, envs, st, _ = docx.import_docx(data)
    ours = sorted("".join(
        r["body"].get("text", "") for r in envs
        if r["kind"] == "run" and r["parent"] == p["oid"])
        for p in envs if p["kind"] == "paragraph")
    ours = sorted(t for t in ours if t.strip())
    assert ours == real_texts, (ours, real_texts)
    assert st["images_imported"] == len(doc.inline_shapes), st

def t_differential_xlsx():
    import datetime
    import openpyxl
    data = (HERE / "samples" / "real-test1.xlsx").read_bytes()
    wb = openpyxl.load_workbook(BytesIO(data))
    _, envs, st, _ = xlsx.import_xlsx(data)
    by_sheet = {}
    for e in envs:
        if e["kind"] == "sheet":
            by_sheet[e["body"]["name"]] = e["body"]["cells"]
    checked = 0
    for ws in wb.worksheets:
        cells = by_sheet[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    assert c.coordinate not in cells or cells[c.coordinate].get("v") is None
                    continue
                v = cells[c.coordinate]["v"]
                if isinstance(c.value, datetime.datetime):
                    serial = (c.value - datetime.datetime(1899, 12, 30)).total_seconds() / 86400
                    assert abs(v - serial) < 1e-6, (c.coordinate, v, c.value)
                elif isinstance(c.value, bool):
                    assert v is c.value, c.coordinate
                elif isinstance(c.value, (int, float)):
                    assert v == c.value or (isinstance(v, float) and abs(v - c.value) < 1e-9), (c.coordinate, v, c.value)
                else:
                    assert v == c.value, (c.coordinate, v, c.value)
                checked += 1
    assert checked == st["cells"] == 96, (checked, st)

def t_differential_pptx():
    from pptx import Presentation
    data = (HERE / "samples" / "real-test.pptx").read_bytes()
    prs = Presentation(BytesIO(data))
    _, envs, st, _ = pptx.import_pptx(data)
    assert st["slides"] == len(prs.slides)
    real_runs = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text:
                            real_runs.append(r.text)
    ours = sorted(r["body"]["text"] for r in envs if r["kind"] == "run")
    assert sorted(real_runs) == ours, (sorted(real_runs), ours)

if __name__ == "__main__":
    print("oracle conformance (real libraries both ways):")
    for name, fn in [("export_docx_readable", t_export_docx_readable),
                     ("export_xlsx_readable", t_export_xlsx_readable),
                     ("export_pptx_readable", t_export_pptx_readable),
                     ("differential_docx", t_differential_docx),
                     ("differential_xlsx", t_differential_xlsx),
                     ("differential_pptx", t_differential_pptx)]:
        check(name, fn)
    print(f"{passed} passed, {failed} failed")
    sys.exit(1 if failed else 0)
